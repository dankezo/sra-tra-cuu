"""EOF JSF crawler. Python 3.10+. See EOF_CRAWLER.md for usage."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import logging
import os
from pathlib import Path
import re
import sqlite3
import threading
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from urllib.parse import urljoin, urlsplit, parse_qs
import xml.etree.ElementTree as ET

import requests
from bs4 import BeautifulSoup

HOME = 'https://services.eof.gr/human-search/home.xhtml'
TABLE = 'frmMain:tblResults'
FIELDS = ['EOF_Code', 'TradeName_Strength', 'Status', 'License_No', 'Procedure', 'Procedure_No', 'Detail_URL']
DETAILS = ['Active_Substance', 'Company_MAH', 'ATC_Code', 'Detail_Status', 'Detail_Error']
LOG = logging.getLogger('eof')


def session():
    s = requests.Session()
    s.headers['User-Agent'] = 'Mozilla/5.0 (compatible; EOF-public-data-export/1.0)'
    return s


def get(s, url):
    for attempt in range(4):
        try:
            r = s.get(url, timeout=(15, 60))
            if r.status_code == 429 or r.status_code >= 500:
                wait = r.headers.get('Retry-After', '')
                time.sleep(min(int(wait), 60) if wait.isdigit() else 2 ** attempt)
                r.raise_for_status()
            r.raise_for_status()
            return r
        except requests.RequestException:
            if attempt == 3:
                raise
            time.sleep(2 ** attempt)


def partial(text):
    root = ET.fromstring(text)
    if root.tag != 'partial-response' or root.find('.//error') is not None or root.find('.//redirect') is not None:
        raise ValueError('JSF session expired, redirected, or returned an error')
    updates = {n.get('id', ''): n.text or '' for n in root.iter('update')}
    state = next((v for k, v in updates.items() if 'javax.faces.ViewState' in k), None)
    if not state:
        raise ValueError('Missing ViewState in AJAX response')
    return updates, state


def form_data(form):
    if form is None:
        raise ValueError('EOF search form not found')
    data = {}
    for el in form.select('input[name], select[name]'):
        if el.has_attr('disabled') or el.get('type') in {'submit', 'button'}:
            continue
        if el.get('type') in {'checkbox', 'radio'} and not el.has_attr('checked'):
            continue
        if el.name == 'select':
            option = el.select_one('option[selected]') or el.select_one('option')
            if option is not None:
                data[el['name']] = option.get('value', option.get_text())
        else:
            data[el['name']] = el.get('value', '')
    return data


def table_rows(html):
    soup = BeautifulSoup(html, 'html.parser')
    body = soup.find('tbody', id=TABLE + '_data')
    out = []
    for tr in (body or soup).find_all('tr'):
        cells = tr.find_all('td', recursive=False)
        if len(cells) < 6:
            continue
        for label in tr.select('.ui-column-title'):
            label.decompose()
        values = [c.get_text(' ', strip=True) for c in cells[:6]]
        link = tr.find('a', href=re.compile(r'view\.xhtml\?'))
        if not values[0] or not link:
            raise ValueError('Product row missing code or detail URL')
        url = urljoin(HOME, link['href'])
        if urlsplit(url).hostname != urlsplit(HOME).hostname:
            raise ValueError('Unexpected detail URL host')
        out.append(dict(zip(FIELDS, values + [url])))
    return out


class Listing:
    def __init__(self):
        self.s = session()
        self.state = ''
        self.total = None

    def post(self, data):
        data['javax.faces.ViewState'] = self.state
        r = self.s.post(HOME, data=data, headers={
            'Faces-Request': 'partial/ajax', 'X-Requested-With': 'XMLHttpRequest',
            'Referer': HOME, 'Origin': 'https://services.eof.gr'}, timeout=(15, 60))
        r.raise_for_status()
        updates, self.state = partial(r.text)
        return updates

    def start(self):
        self.s.close()
        self.s = session()
        soup = BeautifulSoup(get(self.s, HOME).text, 'html.parser')
        data = form_data(soup.find('form', id='frmSearch'))
        self.state = data['javax.faces.ViewState']
        data['frmSearch:txtDrstatus_input'] = 'E'
        data.update({'javax.faces.partial.ajax': 'true', 'javax.faces.source': 'frmSearch:btnSearch',
                     'javax.faces.partial.execute': 'frmSearch', 'javax.faces.partial.render': 'frmMain',
                     'frmSearch:btnSearch': 'frmSearch:btnSearch'})
        updates = self.post(data)
        html = updates.get('frmMain', '')
        match = re.search(r'rowCount\s*:\s*(\d+)', html)
        if not match:
            raise ValueError('Cannot determine total count after search; refusing incomplete export')
        self.total = int(match.group(1))

    def page(self, first):
        for attempt in range(4):
            try:
                data = {'javax.faces.partial.ajax': 'true', 'javax.faces.source': TABLE,
                        'javax.faces.partial.execute': TABLE, 'javax.faces.partial.render': TABLE,
                        'frmMain': 'frmMain', TABLE: TABLE, TABLE + '_pagination': 'true',
                        TABLE + '_first': str(first), TABLE + '_rows': '50',
                        TABLE + '_skipChildren': 'true', TABLE + '_encodeFeature': 'true'}
                updates = self.post(data)
                if TABLE not in updates:
                    raise ValueError('Missing table update')
                return table_rows(updates[TABLE])
            except (requests.RequestException, ValueError, ET.ParseError):
                if attempt == 3:
                    raise
                time.sleep(2 ** attempt)
                self.start()  # Fresh session/state, retry SAME offset; never skip a failed page.


def atomic_csv(path, records):
    tmp = path.with_suffix('.csv.tmp')
    with tmp.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(records)
    os.replace(tmp, path)


def stage1(folder, limit=0, delay=0.4):
    """Replay pagination on resume, validating persisted pages before reusing them."""
    cache = folder / 'pages'
    cache.mkdir(exist_ok=True)
    listing = Listing()
    listing.start()
    total = listing.total
    if not total:
        listing.s.close()
        raise ValueError('Approved search returned zero rows; check form/filter, not a completed crawl')
    target = min(limit, total) if limit else total
    records, codes = [], set()
    scanned = 0
    csv_path = folder / 'eof_stage1.csv'
    first_live = 0
    while first_live < target:
        path = cache / f'{first_live:07d}.json'
        if not path.exists():
            break
        page = json.loads(path.read_text(encoding='utf-8'))
        expected = min(50, total - first_live)
        keys = [r['EOF_Code'] for r in page]
        if len(page) != expected or len(set(keys)) != len(keys) or codes.intersection(keys):
            LOG.warning('Ignoring cache from offset %s', first_live)
            break
        codes.update(keys)
        selected = page[:target - scanned]
        scanned += len(selected)
        records.extend(r for r in selected if norm(r['Status']) in {'εγκεκριμενο', 'approved', 'valid'})
        first_live += 50
    if scanned:
        LOG.info('Stage 1 resume: reused cache through offset %s (%s scanned, %s approved)', first_live - 50, scanned, len(records))
        atomic_csv(csv_path, records)
    try:
        for first in range(first_live, target, 50):
            if first and first % 2000 == 0:
                LOG.info('Refreshing listing session at offset %s', first)
                listing.start()
                if listing.total != total:
                    raise ValueError('EOF count changed mid-run; restart listing')
            page = listing.page(first)
            if listing.total != total:
                raise ValueError('EOF count changed mid-run; restart listing')
            expected = min(50, total - first)
            if len(page) != expected:
                raise ValueError(f'Offset {first}: expected {expected} rows, got {len(page)}')
            keys = [r['EOF_Code'] for r in page]
            if len(set(keys)) != len(keys) or codes.intersection(keys):
                recovered = False
                for attempt in range(4):
                    LOG.warning('Repeated page at offset %s; refresh session (%s/4)', first, attempt + 1)
                    time.sleep(min(2 ** attempt, 8))
                    listing.start()
                    if listing.total != total:
                        raise ValueError('EOF count changed mid-run; restart listing')
                    page = listing.page(first)
                    if listing.total != total:
                        raise ValueError('EOF count changed mid-run; restart listing')
                    if len(page) != expected:
                        continue
                    keys = [r['EOF_Code'] for r in page]
                    if len(set(keys)) == len(keys) and not codes.intersection(keys):
                        recovered = True
                        break
                if not recovered:
                    raise ValueError(f'Repeated page/product at offset {first}; stopped without marking complete')
            codes.update(keys)
            path = cache / f'{first:07d}.json'
            tmp = path.with_suffix('.tmp')
            tmp.write_text(json.dumps(page, ensure_ascii=False), encoding='utf-8')
            os.replace(tmp, path)
            selected = page[:target - scanned]
            scanned += len(selected)
            # Some EOF responses still include special statuses despite the posted filter.
            records.extend(r for r in selected if norm(r['Status']) in {'εγκεκριμενο', 'approved', 'valid'})
            atomic_csv(csv_path, records)
            LOG.info('Stage 1: scanned %s / %s; approved %s', scanned, total, len(records))
            time.sleep(delay)
    finally:
        (folder / 'session_cookies.json').write_text(json.dumps([
            {'name': c.name, 'value': c.value, 'domain': c.domain, 'path': c.path}
            for c in listing.s.cookies]), encoding='utf-8')
        listing.s.close()
    (folder / 'stage1_status.json').write_text(json.dumps({
        'complete': scanned == total, 'rows': len(records), 'scanned': scanned, 'total': total,
        'sha256': hashlib.sha256(csv_path.read_bytes()).hexdigest()}, indent=2), encoding='utf-8')
    return csv_path


def norm(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s.lower()) if not unicodedata.combining(c)).strip(' :')


def detail_fields(html):
    soup = BeautifulSoup(html, 'html.parser')
    if soup.find('form', id='frmSearch'):
        raise ValueError('Detail URL returned search page; rerun phase 1 to renew session-bound URLs')
    result = dict.fromkeys(DETAILS[:3], '')
    for section in soup.select('div.surface-section'):
        heading = section.find('div', recursive=False)
        if not heading:
            continue
        title = norm(heading.get_text(' ', strip=True))
        if title.startswith(('δραστικ', 'active substance')):
            values = []
            for row in section.select('li'):
                cells = row.find_all('div', recursive=False)
                if len(cells) >= 2:
                    values.append(cells[1].get_text(' ', strip=True))
            result['Active_Substance'] = '; '.join(dict.fromkeys(v for v in values if v))
        if 'atc' in title:
            result['ATC_Code'] = '; '.join(dict.fromkeys(re.findall(r'\b[A-Z]\d{2}[A-Z]{2}\d{2}\b', section.get_text(' ', strip=True))))
    labels = {'Active_Substance': ['δραστικ', 'active substance', 'active ingredient'],
              'Company_MAH': ['κατοχος', 'κ.α.κ.', 'mah', 'marketing authorisation holder'],
              'ATC_Code': ['atc', 'ταξινομηση atc', 'κωδικος atc']}
    # Match small label/value containers, never collect whole parent-page text.
    for label in soup.find_all(['label', 'dt', 'th', 'td', 'span', 'div']):
        text = label.get_text(' ', strip=True)
        if len(text) > 100 or not text:
            continue
        key = next((k for k, words in labels.items() if any(norm(text).startswith(w) for w in words)), None)
        if not key:
            continue
        value = ''
        if label.get('for'):
            node = soup.find(id=label['for'])
            if node:
                value = node.get('value', node.get_text(' ', strip=True))
        if not value:
            sibling = label.find_next_sibling()
            if sibling:
                value = sibling.get_text(' ', strip=True)
        if value and len(value) < 2000 and norm(value) != norm(text):
            if key == 'ATC_Code':
                value = '; '.join(dict.fromkeys(re.findall(r'\b[A-Z]\d{2}[A-Z]{2}\d{2}\b', value)))
            if value and not result[key]:
                result[key] = value
    if not any(result.values()):
        raise ValueError('No detail fields found: unexpected HTML/session/login or changed selectors')
    result['Detail_Status'] = 'ok' if all(result.values()) else 'partial'
    result['Detail_Error'] = '' if result['Detail_Status'] == 'ok' else 'Missing: ' + ', '.join(k for k in DETAILS[:3] if not result[k])
    return result


def stage2(folder, workers=6, limit=0, delay=0.3):
    import openpyxl
    path = folder / 'eof_stage1.csv'
    with path.open(encoding='utf-8-sig', newline='') as f:
        records = list(csv.DictReader(f))
    if not records or not set(FIELDS).issubset(records[0]):
        raise ValueError('Empty or invalid stage 1 CSV')
    if limit:
        records = records[:limit]
    db = sqlite3.connect(folder / 'details.sqlite3')
    db.execute('CREATE TABLE IF NOT EXISTS products_v2 (code TEXT PRIMARY KEY, payload TEXT NOT NULL)')
    done = {code: json.loads(data) for code, data in db.execute('SELECT code,payload FROM products_v2')}
    local = threading.local()
    sessions, lock = [], threading.Lock()
    listing = Listing()

    def fetch(url, code, offset):
        if not url or urlsplit(url).hostname != urlsplit(HOME).hostname or not parse_qs(urlsplit(url).query).get('id'):
            return dict.fromkeys(DETAILS[:3], '') | {'Detail_Status': 'error', 'Detail_Error': 'Invalid detail URL'}
        try:
            if not hasattr(local, 'listing'):
                local.listing = Listing()
                local.listing.start()
                local.offset = None
                with lock:
                    sessions.append(local.listing)
            if local.offset != offset:
                local.urls = {r['EOF_Code']: r['Detail_URL'] for r in local.listing.page(offset)}
                local.offset = offset
            url = local.urls[code]
            time.sleep(delay)
            html = get(local.listing.s, url).text
            soup = BeautifulSoup(html, 'html.parser')
            title = soup.title.get_text(' ', strip=True) if soup.title else ''
            if not re.search(r'(?<!\d)' + re.escape(code) + r'(?!\d)', title):
                raise ValueError('Detail product identity mismatch or expired URL; result not accepted')
            return detail_fields(html)
        except Exception as exc:
            return dict.fromkeys(DETAILS[:3], '') | {'Detail_Status': 'error', 'Detail_Error': str(exc)}

    todo = {r['EOF_Code'] for r in records if done.get(r['EOF_Code'], {}).get('Detail_Status') != 'ok'}
    try:
        if todo:
            listing.start()
            with ThreadPoolExecutor(max_workers=workers) as pool:
                for first in range(0, listing.total, 50):
                    # EOF invalidates earlier detail tokens when pagination advances.
                    # Refresh a page, finish its detail batch, THEN advance.
                    page = listing.page(first)
                    futures = {pool.submit(fetch, r['Detail_URL'], r['EOF_Code'], first): r['EOF_Code']
                               for r in page if r['EOF_Code'] in todo}
                    for future in as_completed(futures):
                        code, data = futures[future], future.result()
                        done[code] = data
                        todo.discard(code)
                        db.execute('INSERT OR REPLACE INTO products_v2 VALUES (?, ?)', (code, json.dumps(data, ensure_ascii=False)))
                        db.commit()
                    LOG.info('Stage 2: page offset %s; %s records still to visit', first, len(todo))
                    if not todo:
                        break
                    time.sleep(delay)
    finally:
        db.close()
        listing.s.close()
        for client in sessions:
            client.s.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Medicines'
    ws.append(FIELDS + DETAILS)
    for item in records:
        item = item | done.get(item['EOF_Code'], {'Detail_Status': 'error', 'Detail_Error': 'Product not found in refreshed listing'})
        ws.append([str(item.get(k, '')) for k in FIELDS + DETAILS])
        for cell in ws[ws.max_row]:
            cell.data_type = 's'  # Preserve codes and prevent scraped formula execution.
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    output = folder / 'EOF_Greek_Medicines_Full.xlsx'
    tmp = output.with_suffix('.tmp.xlsx')
    wb.save(tmp)
    os.replace(tmp, output)
    errors = sum(done.get(r['EOF_Code'], {}).get('Detail_Status') != 'ok' for r in records)
    LOG.info('Saved %s rows to %s; %s partial/error rows (retry on next run)', len(records), output, errors)
    return errors


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--phase', choices=['all', '1', '2'], default='all')
    ap.add_argument('--output-dir', type=Path, default=Path('data/raw/GR/crawl'))
    ap.add_argument('--workers', type=int, choices=range(1, 11), default=6)
    ap.add_argument('--limit', type=int, default=0, help='Small test run; 0 means all records')
    ap.add_argument('--delay', type=float, default=0.4)
    args = ap.parse_args()
    if args.limit < 0 or args.delay < 0:
        ap.error('limit and delay must be nonnegative')
    args.output_dir.mkdir(parents=True, exist_ok=True)
    log_file = args.output_dir / 'crawl.log'
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s %(levelname)s %(message)s',
        handlers=[logging.StreamHandler(), logging.FileHandler(log_file, encoding='utf-8')],
        force=True,
    )
    if args.phase in {'all', '1'}:
        status = args.output_dir / 'stage1_status.json'
        status.unlink(missing_ok=True)
        stage1(args.output_dir, args.limit, args.delay)
    if args.phase in {'all', '2'}:
        return 2 if stage2(args.output_dir, args.workers, args.limit, args.delay) else 0
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
