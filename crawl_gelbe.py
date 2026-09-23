"""Gelbe Liste Profi-Suche crawler via ATC level-2 filters.

Python 3.10+. Public catalog research export for DE medicines (excludes unfiltered
/produkte junk by requiring ATC L2). Resume-safe SQLite checkpoints like crawl_eof.py.

Usage:
  python crawl_gelbe.py --phase 1 --delay 0.35
  python crawl_gelbe.py --phase 2 --workers 6
  python crawl_gelbe.py --phase all --limit-atc 2   # smoke: first 2 ATC codes only
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html as html_lib
import json
import logging
import os
import re
import sqlite3
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urljoin, urlsplit

import requests
from bs4 import BeautifulSoup

BASE = "https://www.gelbe-liste.de"
SEARCH = f"{BASE}/profi-suche/results"
FORM = f"{BASE}/profi-suche"
LOG = logging.getLogger("gelbe")

# WHO ATC anatomical/therapeutic level-2 (~90 codes). Empty groups are skipped at runtime.
ATC_L2 = [
    "A01", "A02", "A03", "A04", "A05", "A06", "A07", "A08", "A09", "A10", "A11", "A12", "A13", "A14", "A15", "A16",
    "B01", "B02", "B03", "B05", "B06",
    "C01", "C02", "C03", "C04", "C05", "C07", "C08", "C09", "C10",
    "D01", "D02", "D03", "D04", "D05", "D06", "D07", "D08", "D09", "D10", "D11",
    "G01", "G02", "G03", "G04",
    "H01", "H02", "H03", "H04", "H05",
    "J01", "J02", "J04", "J05", "J06", "J07",
    "L01", "L02", "L03", "L04",
    "M01", "M02", "M03", "M04", "M05", "M09",
    "N01", "N02", "N03", "N04", "N05", "N06", "N07",
    "P01", "P02", "P03",
    "R01", "R02", "R03", "R05", "R06", "R07",
    "S01", "S02", "S03",
    "V01", "V03", "V04", "V06", "V07", "V08", "V09", "V10", "V20",
]

LIST_FIELDS = ["Product_ID", "Drug_Name", "Company", "Detail_URL", "ATC_Search_Groups"]
DETAIL_FIELDS = [
    "Active_Substance", "Full_ATC", "Dosage_Form", "Strength", "PZN",
    "Detail_Status", "Detail_Error",
]
PRODUCT_ID_RE = re.compile(r"/produkte/[^\"'?#]+_(\d+)\b", re.I)
ATC_FULL_RE = re.compile(r"\b([A-Z]\d{2}[A-Z]{2}\d{2})\b")


def session() -> requests.Session:
    s = requests.Session()
    # Site returns 410 for non-browser UAs; use a normal desktop Chrome header.
    s.headers.update({
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        ),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "de-DE,de;q=0.9,en;q=0.8",
    })
    return s


def get(s: requests.Session, url: str, **kwargs):
    for attempt in range(5):
        try:
            r = s.get(url, timeout=(15, 60), **kwargs)
            if r.status_code in {429, 503} or r.status_code >= 500:
                wait = r.headers.get("Retry-After", "")
                time.sleep(min(int(wait), 60) if wait.isdigit() else 2 ** attempt)
                continue
            r.raise_for_status()
            return r
        except requests.RequestException:
            if attempt == 4:
                raise
            time.sleep(2 ** attempt)
    raise RuntimeError(f"GET failed: {url}")


def post(s: requests.Session, url: str, data: dict):
    for attempt in range(5):
        try:
            r = s.post(url, data=data, timeout=(15, 60))
            if r.status_code in {429, 503} or r.status_code >= 500:
                wait = r.headers.get("Retry-After", "")
                time.sleep(min(int(wait), 60) if wait.isdigit() else 2 ** attempt)
                continue
            r.raise_for_status()
            return r
        except requests.RequestException:
            if attempt == 4:
                raise
            time.sleep(2 ** attempt)
    raise RuntimeError(f"POST failed: {url}")


def search_params(atc: str, page: int | None = None) -> dict:
    data = {
        "product": "",
        "company": "",
        "substance": "",
        "atc": atc,
        "icd": "",
        "dispensing": "100",
        "divisible": "false",
        "crushable": "false",
        "watersoluble": "false",
        "hasPicture": "false",
        "hasSPC": "false",
    }
    if page and page > 1:
        data["page"] = str(page)
    return data


def parse_count(html: str) -> int | None:
    text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
    m = re.search(r"([\d.]+)\s*Pr[aä]parate", text, re.I)
    if not m:
        return None
    return int(m.group(1).replace(".", ""))


def parse_list_page(html: str) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    ul = soup.find("ul", class_="product-list")
    if not ul:
        return []
    out = []
    for li in ul.find_all("li", recursive=False):
        a = li.find("a", href=True)
        if not a:
            continue
        href = a["href"].strip()
        m = PRODUCT_ID_RE.search(href)
        if not m:
            continue
        pid = m.group(1)
        name_tag = a.find("h5") or a.find("span", class_=re.compile(r"text-link|font-weight"))
        company_tag = a.find("p", class_="small") or a.find("p")
        name = name_tag.get_text(" ", strip=True) if name_tag else ""
        company = company_tag.get_text(" ", strip=True) if company_tag else ""
        name = html_lib.unescape(name)
        company = html_lib.unescape(company)
        url = urljoin(BASE, href)
        if urlsplit(url).hostname != "www.gelbe-liste.de":
            continue
        out.append({
            "Product_ID": pid,
            "Drug_Name": name,
            "Company": company,
            "Detail_URL": url,
        })
    return out


def has_next_page(html: str) -> bool:
    soup = BeautifulSoup(html, "html.parser")
    return bool(soup.select_one('link[rel="next"], a[rel="next"]'))


def open_db(folder: Path) -> sqlite3.Connection:
    db = sqlite3.connect(folder / "gelbe.sqlite3", timeout=60)
    db.execute("PRAGMA journal_mode=WAL")
    db.execute(
        """CREATE TABLE IF NOT EXISTS products (
            product_id TEXT PRIMARY KEY,
            drug_name TEXT NOT NULL,
            company TEXT NOT NULL,
            detail_url TEXT NOT NULL,
            atc_search_groups TEXT NOT NULL DEFAULT '',
            updated_at REAL NOT NULL
        )"""
    )
    db.execute(
        """CREATE TABLE IF NOT EXISTS atc_progress (
            atc TEXT PRIMARY KEY,
            expected_count INTEGER,
            pages_done INTEGER NOT NULL DEFAULT 0,
            items_seen INTEGER NOT NULL DEFAULT 0,
            complete INTEGER NOT NULL DEFAULT 0,
            updated_at REAL NOT NULL
        )"""
    )
    db.execute(
        """CREATE TABLE IF NOT EXISTS details (
            product_id TEXT PRIMARY KEY,
            payload TEXT NOT NULL,
            status TEXT NOT NULL,
            updated_at REAL NOT NULL
        )"""
    )
    db.commit()
    return db


def upsert_product(db: sqlite3.Connection, item: dict, atc: str) -> None:
    row = db.execute(
        "SELECT atc_search_groups FROM products WHERE product_id=?",
        (item["Product_ID"],),
    ).fetchone()
    groups = set()
    if row and row[0]:
        groups.update(g for g in row[0].split("|") if g)
    groups.add(atc)
    db.execute(
        """INSERT INTO products(product_id, drug_name, company, detail_url, atc_search_groups, updated_at)
           VALUES(?,?,?,?,?,?)
           ON CONFLICT(product_id) DO UPDATE SET
             drug_name=excluded.drug_name,
             company=excluded.company,
             detail_url=excluded.detail_url,
             atc_search_groups=excluded.atc_search_groups,
             updated_at=excluded.updated_at""",
        (
            item["Product_ID"],
            item["Drug_Name"],
            item["Company"],
            item["Detail_URL"],
            "|".join(sorted(groups)),
            time.time(),
        ),
    )


def export_stage1_csv(db: sqlite3.Connection, path: Path) -> int:
    rows = list(db.execute(
        "SELECT product_id, drug_name, company, detail_url, atc_search_groups FROM products ORDER BY product_id"
    ))
    tmp = path.with_suffix(".csv.tmp")
    with tmp.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=LIST_FIELDS)
        w.writeheader()
        for pid, name, company, url, groups in rows:
            w.writerow({
                "Product_ID": pid,
                "Drug_Name": name,
                "Company": company,
                "Detail_URL": url,
                "ATC_Search_Groups": groups,
            })
    os.replace(tmp, path)
    return len(rows)


def stage1(folder: Path, delay: float = 0.35, limit_atc: int = 0, only_atc: list[str] | None = None) -> Path:
    codes = only_atc or ATC_L2
    if limit_atc:
        codes = codes[:limit_atc]
    db = open_db(folder)
    s = session()
    get(s, FORM)  # warm cookies
    try:
        for i, atc in enumerate(codes, 1):
            prog = db.execute(
                "SELECT pages_done, complete, expected_count FROM atc_progress WHERE atc=?",
                (atc,),
            ).fetchone()
            if prog and prog[1]:
                LOG.info("Stage 1 [%s/%s] %s already complete — skip", i, len(codes), atc)
                continue

            start_page = (prog[0] + 1) if prog and prog[0] else 1
            LOG.info("Stage 1 [%s/%s] ATC %s from page %s", i, len(codes), atc, start_page)

            # Page 1 via POST (matches HAR); later pages via GET with params.
            if start_page == 1:
                r = post(s, SEARCH, search_params(atc))
            else:
                r = get(s, SEARCH, params=search_params(atc, start_page))
            expected = parse_count(r.text)
            items = parse_list_page(r.text)
            if expected == 0 or (expected is None and not items):
                db.execute(
                    """INSERT INTO atc_progress(atc, expected_count, pages_done, items_seen, complete, updated_at)
                       VALUES(?,?,?,?,?,?)
                       ON CONFLICT(atc) DO UPDATE SET
                         expected_count=excluded.expected_count, pages_done=excluded.pages_done,
                         items_seen=excluded.items_seen, complete=1, updated_at=excluded.updated_at""",
                    (atc, 0, 0, 0, 1, time.time()),
                )
                db.commit()
                LOG.info("ATC %s: empty", atc)
                time.sleep(delay)
                continue

            page = start_page
            if start_page == 1:
                seen_on_atc = 0
            else:
                prev = db.execute(
                    "SELECT items_seen FROM atc_progress WHERE atc=?", (atc,)
                ).fetchone()
                seen_on_atc = prev[0] if prev else 0

            while True:
                if page > start_page:
                    r = get(s, SEARCH, params=search_params(atc, page))
                    items = parse_list_page(r.text)
                    if expected is None:
                        expected = parse_count(r.text)

                if not items:
                    LOG.warning("ATC %s page %s empty — stop group", atc, page)
                    break

                for item in items:
                    upsert_product(db, item, atc)
                seen_on_atc += len(items)
                db.execute(
                    """INSERT INTO atc_progress(atc, expected_count, pages_done, items_seen, complete, updated_at)
                       VALUES(?,?,?,?,0,?)
                       ON CONFLICT(atc) DO UPDATE SET
                         expected_count=COALESCE(excluded.expected_count, atc_progress.expected_count),
                         pages_done=excluded.pages_done,
                         items_seen=excluded.items_seen,
                         updated_at=excluded.updated_at""",
                    (atc, expected, page, seen_on_atc, time.time()),
                )
                db.commit()
                LOG.info(
                    "ATC %s page %s: +%s items (group seen %s%s)",
                    atc, page, len(items), seen_on_atc,
                    f" / {expected}" if expected else "",
                )

                if not has_next_page(r.text):
                    break
                page += 1
                time.sleep(delay)

            db.execute(
                "UPDATE atc_progress SET complete=1, updated_at=? WHERE atc=?",
                (time.time(), atc),
            )
            db.commit()
            n_unique = db.execute("SELECT COUNT(*) FROM products").fetchone()[0]
            LOG.info("ATC %s done. Unique products so far: %s", atc, n_unique)
            export_stage1_csv(db, folder / "gelbe_stage1.csv")
            time.sleep(delay)

            # Refresh session cookies periodically
            if i % 15 == 0:
                get(s, FORM)
    finally:
        csv_path = folder / "gelbe_stage1.csv"
        n = export_stage1_csv(db, csv_path)
        complete = db.execute("SELECT COUNT(*) FROM atc_progress WHERE complete=1").fetchone()[0]
        status = {
            "complete": complete >= len(codes),
            "atc_done": complete,
            "atc_total": len(codes),
            "unique_products": n,
            "sha256": hashlib.sha256(csv_path.read_bytes()).hexdigest() if csv_path.exists() else "",
        }
        (folder / "stage1_status.json").write_text(json.dumps(status, indent=2), encoding="utf-8")
        db.close()
        s.close()
        LOG.info("Stage 1 finished: %s unique products (%s/%s ATC complete)", n, complete, len(codes))
    return folder / "gelbe_stage1.csv"


def extract_product_data(html: str) -> dict | None:
    """Parse the HTML-entity-encoded ProductData JSON embedded in detail pages."""
    m = re.search(r"ProductData\"\s*:\s*(\{&quot;Product&quot;:.+)", html)
    if not m:
        m = re.search(r"ProductData\"\s*:\s*(\{\"Product\":.+)", html)
    if not m:
        return None
    raw = m.group(1)
    un = html_lib.unescape(raw) if "&quot;" in raw[:40] else raw
    start = un.find("{")
    if start < 0:
        return None
    depth = 0
    end = None
    in_str = False
    esc = False
    for i, ch in enumerate(un[start:], start):
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
            continue
        if ch == '"':
            in_str = True
        elif ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                end = i + 1
                break
    if end is None:
        return None
    try:
        data = json.loads(un[start:end])
    except json.JSONDecodeError:
        return None
    return data.get("Product") if isinstance(data, dict) else None


def parse_detail(html: str, product_id: str) -> dict:
    soup = BeautifulSoup(html, "html.parser")
    out = {k: "" for k in DETAIL_FIELDS}
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    if "nicht gefunden" in title.lower() or title.strip() == "404":
        out["Detail_Status"] = "error"
        out["Detail_Error"] = "Product not found"
        return out

    product = extract_product_data(html)
    if product:
        items = product.get("ItemList") or []
        item = items[0] if items else {}
        # Active ingredients: MoleculetypeCode "A"
        actives = []
        strengths = []
        for el in item.get("COMPOSITIONELEMENTS_LIST") or []:
            if str(el.get("MoleculetypeCode", "")).upper() != "A":
                continue
            name = (el.get("MoleculeName") or "").strip()
            if name:
                actives.append(name)
            mass = el.get("MassFrom")
            unit = (el.get("MoleculeUnitName") or "").strip()
            if mass is not None and name:
                # MassFrom sometimes stored scaled (e.g. 100000 I.E.)
                strengths.append(f"{mass} {unit}".strip() + f" {name}")
        if not actives:
            for atc in item.get("ATCCODE_LIST") or []:
                n = (atc.get("Name") or "").strip()
                if n:
                    actives.append(n)
        out["Active_Substance"] = "; ".join(dict.fromkeys(actives))
        atcs = [a.get("Code") for a in (item.get("ATCCODE_LIST") or []) if a.get("Code")]
        if not atcs:
            atcs = ATC_FULL_RE.findall(json.dumps(product, ensure_ascii=False))
        out["Full_ATC"] = "; ".join(dict.fromkeys(atcs))
        out["Dosage_Form"] = (item.get("PharmformName") or item.get("Name") or "").strip()
        if strengths:
            out["Strength"] = "; ".join(strengths[:3])
        elif product.get("Name"):
            sm = re.search(
                r"(\d+[.,]?\d*\s*(?:mg|g|µg|mcg|ml|I\.?E\.?|E\.?|%)(?:\s*/\s*\d+[.,]?\d*\s*(?:ml|g|mg))?)",
                product["Name"],
                re.I,
            )
            if sm:
                out["Strength"] = sm.group(1).strip()
        # Optional PZN not always in ProductData; fall through to JSON-LD below.

    # Fallbacks when ProductData missing/partial
    if not out["Active_Substance"]:
        substances = []
        for a in soup.find_all("a", href=True):
            href = a["href"]
            if "/wirkstoffe/" in href and not href.rstrip("/").endswith("wirkstoffe"):
                t = a.get_text(" ", strip=True)
                if t:
                    substances.append(t)
        if not substances:
            meta = soup.find("meta", attrs={"name": "description"})
            if meta and meta.get("content"):
                m = re.search(r"Wirkstoff\s+(.+?)\s*\(ATC", meta["content"], re.I)
                if m:
                    substances.append(m.group(1).strip())
            kw = soup.find("meta", attrs={"name": "keywords"})
            if kw and kw.get("content"):
                parts = [p.strip() for p in kw["content"].split(",")]
                # keywords: name, substance, ATC, code, substance, company
                if len(parts) >= 2:
                    substances.append(parts[1])
        out["Active_Substance"] = "; ".join(dict.fromkeys(s for s in substances if s))

    if not out["Full_ATC"]:
        labeled = []
        text = soup.get_text(" ", strip=True)
        for m in re.finditer(r"ATC[-\s]?Code[^A-Z]{0,30}([A-Z]\d{2}[A-Z]{2}\d{2})", text, re.I):
            labeled.append(m.group(1))
        meta = soup.find("meta", attrs={"name": "keywords"})
        if meta and meta.get("content"):
            labeled.extend(ATC_FULL_RE.findall(meta["content"]))
        m = re.search(r"product_ATC_class_code['\"]?\s*[:=]\s*\[?['\"]([A-Z]\d{2}[A-Z]{2}\d{2})", html)
        if m:
            labeled.append(m.group(1))
        out["Full_ATC"] = "; ".join(dict.fromkeys(labeled or ATC_FULL_RE.findall(html)))

    if not out["Dosage_Form"]:
        for label in ("Darreichungsform", "Dosage form", "Arzneiform"):
            node = soup.find(string=re.compile(re.escape(label), re.I))
            if not node:
                continue
            parent = node.parent
            sib = parent.find_next_sibling() if parent else None
            if sib:
                out["Dosage_Form"] = sib.get_text(" ", strip=True)
                break

    if not out["Strength"]:
        h1 = soup.find("h1")
        name = h1.get_text(" ", strip=True) if h1 else ""
        sm = re.search(
            r"(\d+[.,]?\d*\s*(?:mg|g|µg|mcg|ml|I\.?E\.?|E\.?|%)(?:\s*/\s*\d+[.,]?\d*\s*(?:ml|g|mg))?)",
            name,
            re.I,
        )
        if sm:
            out["Strength"] = sm.group(1).strip()

    if not out["PZN"]:
        for script in soup.find_all("script", attrs={"type": "application/ld+json"}):
            raw = script.string or script.get_text() or ""
            try:
                data = json.loads(raw)
            except json.JSONDecodeError:
                continue
            items = data if isinstance(data, list) else [data]
            for obj in items:
                if not isinstance(obj, dict):
                    continue
                for key in ("sku", "productID", "gtin13", "gtin"):
                    if obj.get(key):
                        out["PZN"] = str(obj[key])
                        break

    missing = [k for k in ("Active_Substance", "Full_ATC") if not out[k]]
    if not missing:
        out["Detail_Status"] = "ok"
        out["Detail_Error"] = ""
    elif out["Active_Substance"] or out["Full_ATC"] or out["Dosage_Form"]:
        out["Detail_Status"] = "partial"
        out["Detail_Error"] = "Missing: " + ", ".join(missing)
    else:
        out["Detail_Status"] = "error"
        out["Detail_Error"] = "No detail fields found"
    return out


def stage2(folder: Path, workers: int = 6, delay: float = 0.3, limit: int = 0) -> int:
    db = open_db(folder)
    rows = list(db.execute(
        "SELECT product_id, drug_name, company, detail_url, atc_search_groups FROM products ORDER BY product_id"
    ))
    if not rows:
        raise ValueError("No products in SQLite — run phase 1 first")
    if limit:
        rows = rows[:limit]

    done = {
        pid: json.loads(payload)
        for pid, payload, status in db.execute("SELECT product_id, payload, status FROM details")
        if status == "ok"
    }
    todo = [r for r in rows if r[0] not in done]
    LOG.info("Stage 2: %s products; %s already ok; %s to fetch", len(rows), len(done), len(todo))

    local = threading.local()
    lock = threading.Lock()
    sessions: list[requests.Session] = []

    def client() -> requests.Session:
        if not hasattr(local, "s"):
            local.s = session()
            get(local.s, FORM)
            with lock:
                sessions.append(local.s)
        return local.s

    def fetch(row):
        pid, name, company, url, groups = row
        try:
            time.sleep(delay)
            html = get(client(), url).text
            data = parse_detail(html, pid)
            return pid, data
        except Exception as exc:
            return pid, {
                "Active_Substance": "", "Full_ATC": "", "Dosage_Form": "", "Strength": "", "PZN": "",
                "Detail_Status": "error", "Detail_Error": str(exc),
            }

    try:
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {pool.submit(fetch, row): row[0] for row in todo}
            finished = 0
            for fut in as_completed(futures):
                pid, data = fut.result()
                status = data.get("Detail_Status", "error")
                db.execute(
                    """INSERT INTO details(product_id, payload, status, updated_at) VALUES(?,?,?,?)
                       ON CONFLICT(product_id) DO UPDATE SET
                         payload=excluded.payload, status=excluded.status, updated_at=excluded.updated_at""",
                    (pid, json.dumps(data, ensure_ascii=False), status, time.time()),
                )
                if finished % 25 == 0:
                    db.commit()
                finished += 1
                if finished % 50 == 0 or finished == len(todo):
                    db.commit()
                    LOG.info("Stage 2: %s / %s detail fetches", finished, len(todo))
        db.commit()
    finally:
        for s in sessions:
            s.close()

    # Merge all details for export
    all_details = {
        pid: json.loads(payload)
        for pid, payload in db.execute("SELECT product_id, payload FROM details")
    }

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicines"
    headers = LIST_FIELDS + DETAIL_FIELDS
    ws.append(headers)
    errors = 0
    for pid, name, company, url, groups in rows:
        detail = all_details.get(pid, {
            "Active_Substance": "", "Full_ATC": "", "Dosage_Form": "", "Strength": "", "PZN": "",
            "Detail_Status": "error", "Detail_Error": "Not fetched",
        })
        if detail.get("Detail_Status") != "ok":
            errors += 1
        row = {
            "Product_ID": pid,
            "Drug_Name": name,
            "Company": company,
            "Detail_URL": url,
            "ATC_Search_Groups": groups,
            **detail,
        }
        ws.append([str(row.get(k, "")) for k in headers])
        for cell in ws[ws.max_row]:
            cell.data_type = "s"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = folder / "GelbeListe_Medicines_ATC_Filtered.xlsx"
    tmp = out.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, out)
    db.close()
    LOG.info("Saved %s rows to %s; %s partial/error (retry phase 2)", len(rows), out, errors)
    return errors


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--phase", choices=["all", "1", "2"], default="all")
    ap.add_argument("--output-dir", type=Path, default=Path("data/raw/DE/crawl"))
    ap.add_argument("--workers", type=int, default=6)
    ap.add_argument("--delay", type=float, default=0.35)
    ap.add_argument("--limit", type=int, default=0, help="Limit detail rows in phase 2 (0=all)")
    ap.add_argument("--limit-atc", type=int, default=0, help="Only first N ATC codes (smoke test)")
    ap.add_argument("--atc", action="append", default=[], help="Restrict to these ATC L2 codes (repeatable)")
    args = ap.parse_args()
    if args.workers < 1 or args.workers > 10:
        ap.error("workers must be 1..10")
    if args.delay < 0 or args.limit < 0 or args.limit_atc < 0:
        ap.error("delay/limit/limit-atc must be nonnegative")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.StreamHandler(),
            logging.FileHandler(args.output_dir / "crawl.log", encoding="utf-8"),
        ],
        force=True,
    )
    only = [c.upper() for c in args.atc] or None
    if args.phase in {"all", "1"}:
        stage1(args.output_dir, args.delay, args.limit_atc, only)
    if args.phase in {"all", "2"}:
        return 2 if stage2(args.output_dir, args.workers, args.delay, args.limit) else 0
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
